from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "report"
IMAGE_PATH = ROOT / "images" / "dashboard_overview.png"
OUTPUT_PATH = REPORT_DIR / "AI_Supplier_Risk_Project_Report.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_bullets(document, items):
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.add_run(item)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(28)
styles["Title"].font.bold = True
for style_name, size in [("Heading 1", 18), ("Heading 2", 13)]:
    styles[style_name].font.name = "Aptos Display"
    styles[style_name].font.size = Pt(size)
    styles[style_name].font.color.rgb = RGBColor(31, 78, 121)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("AI-Assisted Supplier Performance\n& Risk Monitoring")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Procurement Analytics Portfolio Project")
run.bold = True
run.font.size = Pt(15)
run.font.color.rgb = RGBColor(31, 78, 121)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Python • pandas • scikit-learn • Power BI • DAX\n").italic = True
meta.add_run("Simulated data | Human-in-the-loop decision support")

doc.add_heading("Executive Summary", level=1)
doc.add_paragraph(
    "This project demonstrates how a procurement team can combine supplier-performance "
    "KPIs with machine-learning-assisted anomaly detection to prioritize supplier reviews. "
    "Using simulated purchase-order data, the analysis evaluates 106 suppliers across "
    "delivery, quality, pricing, and contract-compliance dimensions. A transparent business "
    "risk score is combined with an Isolation Forest anomaly signal and presented in an "
    "interactive Power BI dashboard."
)

kpi_table = doc.add_table(rows=2, cols=4)
kpi_table.alignment = WD_TABLE_ALIGNMENT.CENTER
kpi_table.style = "Table Grid"
for i, label in enumerate(["Suppliers", "High Risk", "AI Review", "Simulated Spend"]):
    cell = kpi_table.cell(0, i)
    cell.text = label
    set_cell_shading(cell, "D9EAF7")
    cell.paragraphs[0].runs[0].bold = True
for i, value in enumerate(["106", "6", "7", "411.18 million"]):
    cell = kpi_table.cell(1, i)
    cell.text = value
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.paragraphs[0].runs[0].bold = True

doc.add_page_break()
doc.add_heading("1. Business Context", level=1)
doc.add_paragraph(
    "Procurement teams need to monitor supplier performance before delivery and quality "
    "problems create material shortages, production disruption, or unplanned cost. This "
    "project converts detailed purchase-order data into a practical supplier-review queue."
)
add_bullets(doc, [
    "Which suppliers have the highest operational risk?",
    "What factors are driving each supplier's risk?",
    "Which suppliers should procurement review first?",
    "What corrective actions should the procurement team consider?",
])

doc.add_heading("2. Data and KPI Framework", level=1)
doc.add_paragraph(
    "The simulated source table contains purchase-order dates, promised and actual delivery "
    "dates, supplier identifiers, country and category, spend, contract status, and quality-"
    "rejection indicators. The order-level records are aggregated into one supplier-level scorecard."
)

kpis = [
    ("On-time rate", "Share of orders delivered on or before the promised date"),
    ("Late rate", "Share of orders delivered after the promised date"),
    ("Average delay days", "Average number of late days across a supplier's orders"),
    ("Reject rate", "Share of orders with a quality rejection"),
    ("Price volatility", "Variation in unit price over time"),
    ("Off-contract share", "Share of purchase orders placed outside contract"),
]
table = doc.add_table(rows=1, cols=2)
table.style = "Light Shading Accent 1"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.rows[0].cells[0].text = "KPI"
table.rows[0].cells[1].text = "Business meaning"
for label, meaning in kpis:
    cells = table.add_row().cells
    cells[0].text = label
    cells[1].text = meaning

doc.add_heading("3. Methodology", level=1)
doc.add_heading("3.1 Transparent business risk score", level=2)
doc.add_paragraph(
    "Higher late-delivery rates, longer average delays, higher reject rates, greater price "
    "volatility, and more off-contract purchases increase the business risk score. The score "
    "is explainable because its drivers can be traced to operational KPIs."
)
doc.add_heading("3.2 AI-assisted anomaly detection", level=2)
doc.add_paragraph(
    "Isolation Forest highlights suppliers whose combined KPI pattern differs from the broader "
    "supplier population. The anomaly flag is a review signal, not a prediction of supplier "
    "failure and not an autonomous sourcing decision."
)
doc.add_heading("3.3 Combined review priority", level=2)
doc.add_paragraph(
    "The final risk score combines the business risk score with the anomaly signal. Suppliers "
    "are grouped into Low, Medium, and High tiers to create a practical, human-reviewed queue."
)

doc.add_page_break()
doc.add_heading("4. Dashboard and Results", level=1)
if IMAGE_PATH.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(IMAGE_PATH), width=Inches(7.0))
    caption = doc.add_paragraph("Figure 1. Power BI supplier risk overview")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True

doc.add_paragraph(
    "The dashboard identifies 6 High Risk suppliers and flags 7 suppliers for AI-assisted "
    "review. The supplier population consists of 78 Low, 22 Medium, and 6 High risk suppliers."
)

ranked = [
    ("1", "Great Wall Packaging", "95.4"),
    ("2", "Fujikawa Resources", "93.8"),
    ("3", "Thames Industrial Supply", "93.0"),
    ("4", "Cascade Technologies", "92.6"),
    ("5", "Great Wall Components", "92.6"),
    ("6", "Baltic Workplace", "91.7"),
]
table = doc.add_table(rows=1, cols=3)
table.style = "Light Shading Accent 1"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, text_value in zip(table.rows[0].cells, ["Rank", "Supplier", "Final risk score"]):
    cell.text = text_value
for row in ranked:
    cells = table.add_row().cells
    for cell, text_value in zip(cells, row):
        cell.text = text_value

doc.add_heading("5. Procurement Recommendations", level=1)
add_bullets(doc, [
    "Start supplier recovery plans with the six High Risk suppliers, beginning with the top three.",
    "Review delivery and quality KPIs weekly until performance remains within agreed thresholds.",
    "Require root-cause analysis and corrective-action evidence for repeated late delivery or rejection.",
    "Check whether high-risk suppliers support critical or single-source items and prepare alternatives.",
    "Review anomaly flags together with purchase-order volume and operational context before escalation.",
])

doc.add_heading("6. Governance and Limitations", level=1)
add_bullets(doc, [
    "The dataset is simulated and is suitable for portfolio demonstration, not real supplier decisions.",
    "Risk weights and thresholds are business assumptions that require stakeholder validation.",
    "An anomaly is an unusual pattern, not evidence of wrongdoing or inevitable failure.",
    "Suppliers with small order counts may show unstable percentages.",
    "The dashboard supports prioritization; it does not replace buyer judgment or supplier communication.",
    "Spend currency is intentionally unspecified and must not be represented as real company financials.",
])

doc.add_heading("7. Technical Stack and Deliverables", level=1)
add_bullets(doc, [
    "Python and pandas for data preparation and aggregation",
    "scikit-learn Isolation Forest for anomaly detection",
    "Power BI and DAX for interactive reporting",
    "CSV outputs for transparent handoff between analytics and visualization",
    "PBIX dashboard, exported dashboard PDF, source code, datasets, README, and interview notes",
])

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("AI-Assisted Supplier Performance & Risk Monitoring | Portfolio Project")

REPORT_DIR.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT_PATH)
print(OUTPUT_PATH)
